"""Motor da automação: lê relatórios da Stripe, consolida, confere e gera o Excel."""
import re
import unicodedata
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FONTE = "Arial"

RENOMEAR = {
    "automatic_payout_id": "id_repasse",
    "automatic_payout_effective_at": "data_repasse",
    "balance_transaction_id": "id_transacao",
    "created": "data_transacao",
    "currency": "moeda",
    "gross": "bruto",
    "fee": "tarifa",
    "net": "liquido",
    "reporting_category": "categoria",
    "source_id": "id_origem",
    "description": "descricao",
    "customer_id": "id_cliente",
    "customer_email": "email_cliente",
    # --- exportação "balance history" (inglês) ---
    "id": "id_transacao",
    "type": "categoria",
    "amount": "bruto",
    "fees": "tarifa",
    "customer id": "id_cliente",
    "customer email": "email_cliente",
    "customer name": "cliente",
    "converted amount": "valor_convertido",
    "converted currency": "moeda_convertida",
    "details": "detalhes",
    # --- exportações em português ---
    "data / hora": "data_transacao",
    "data/hora": "data_transacao",
    "data": "data_transacao",
    "tipo": "categoria",
    "id transacao": "id_transacao",
    "id da transacao": "id_transacao",
    "cliente": "cliente",
    "id do cliente": "id_cliente",
    "e-mail do cliente": "email_cliente",
    "descricao": "descricao",
    "moeda": "moeda",
    "bruto": "bruto",
    "valor bruto": "bruto",
    "valor": "bruto",
    "tarifa": "tarifa",
    "tarifas": "tarifa",
    "taxa": "tarifa",
    "liquido": "liquido",
    "valor liquido": "liquido",
    "id do repasse": "id_repasse",
    "data do repasse": "data_repasse",
    "id de origem": "id_origem",
}
ORDEM = [
    "arquivo_origem", "id_repasse", "data_repasse", "id_transacao", "data_transacao",
    "cliente", "id_cliente", "email_cliente", "cpf_cnpj", "categoria", "moeda",
    "bruto", "tarifa", "liquido", "valor_convertido", "moeda_convertida",
    "id_origem", "descricao", "detalhes",
]
OBRIGATORIAS = ["id_transacao"]
COLUNAS_VALOR = ["bruto", "tarifa", "liquido"]
COLUNAS_DATA = ["data_repasse", "data_transacao"]
PALAVRAS_DOC = ["cpf", "cnpj", "tax_id", "tax id", "documento"]
EXTENSOES = (".csv", ".xlsx", ".xls")

COR_CAB = PatternFill("solid", start_color="1F3864")
FONTE_CAB = Font(name=FONTE, bold=True, color="FFFFFF")


# ------------------------------------------------------------------
# Leitura
# ------------------------------------------------------------------
def _norm(nome) -> str:
    """'Líquido ' -> 'liquido'; 'CNPJ / CPF (Metadata)' -> 'cnpj / cpf (metadata)'."""
    t = unicodedata.normalize("NFKD", str(nome)).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", t).strip().lower()


def _pontuacao(linha) -> int:
    """Quantas células da linha parecem nomes de coluna conhecidos."""
    pts = 0
    for v in linha:
        if pd.isna(v):
            continue
        n = _norm(v)
        if n in RENOMEAR or n in RENOMEAR.values() or any(p in n for p in PALAVRAS_DOC):
            pts += 1
    return pts


def _achar_cabecalho(bruto: pd.DataFrame):
    """Devolve (indice_da_linha, pontos) da melhor candidata a cabeçalho nas 30 primeiras linhas."""
    melhor = (0, 0)
    for i in range(min(30, len(bruto))):
        pts = _pontuacao(bruto.iloc[i].tolist())
        if pts > melhor[1]:
            melhor = (i, pts)
    return melhor


def _aplicar_cabecalho(bruto: pd.DataFrame, idx: int) -> pd.DataFrame:
    nomes = []
    for j, v in enumerate(bruto.iloc[idx].tolist()):
        n = "" if pd.isna(v) else _norm(v)
        nomes.append(n or f"coluna_{j + 1}")
    df = bruto.iloc[idx + 1:].copy()
    df.columns = nomes
    df = df.dropna(how="all").dropna(axis=1, how="all")
    return df.reset_index(drop=True)


def ler_tabela(caminho: Path) -> pd.DataFrame:
    """Lê CSV/Excel mesmo com linhas de título no topo ou várias abas."""
    caminho = Path(caminho)
    if caminho.suffix.lower() == ".csv":
        try:
            abas = {"csv": pd.read_csv(caminho, dtype=str, sep=None, engine="python",
                                       encoding="utf-8-sig", header=None)}
        except Exception:
            abas = {"csv": pd.read_csv(caminho, dtype=str, sep=None, engine="python",
                                       encoding="utf-8-sig")}
            abas["csv"] = abas["csv"].T.reset_index().T.reset_index(drop=True)
    else:
        abas = pd.read_excel(caminho, dtype=str, header=None, sheet_name=None)

    melhor = None  # (pontos, nome_aba, bruto, idx)
    for nome, bruto in abas.items():
        idx, pts = _achar_cabecalho(bruto)
        if melhor is None or pts > melhor[0]:
            melhor = (pts, nome, bruto, idx)
    pts, nome, bruto, idx = melhor
    return _aplicar_cabecalho(bruto, idx)


def ler_repasses(arquivos, log=print) -> pd.DataFrame:
    partes = []
    for arq in arquivos:
        arq = Path(arq)
        df = ler_tabela(arq)
        df["arquivo_origem"] = arq.name
        partes.append(df)
        log(f"  {arq.name}: {len(df)} linhas")
    return pd.concat(partes, ignore_index=True)


# ------------------------------------------------------------------
# CPF / CNPJ
# ------------------------------------------------------------------
def extrair_documento(df: pd.DataFrame) -> pd.DataFrame:
    cols = [c for c in df.columns if any(p in c for p in PALAVRAS_DOC)]
    if cols:
        limpo = df[cols].replace(r"^\s*$", np.nan, regex=True)
        df["cpf_cnpj"] = limpo.bfill(axis=1).iloc[:, 0]
    else:
        df["cpf_cnpj"] = np.nan
    return df


def completar_com_cadastro(df: pd.DataFrame, arquivos_cadastro, log=print) -> pd.DataFrame:
    if not arquivos_cadastro:
        return df
    chave = "id_cliente" if "id_cliente" in df.columns else ("cliente" if "cliente" in df.columns else None)
    if chave is None:
        log("  AVISO: o relatório não tem coluna de cliente; o cadastro não pôde ser cruzado.")
        return df

    candidatas = ("id", "customer_id", "id_cliente", "id do cliente", "cliente", "e-mail", "email")
    partes = []
    for arq in arquivos_cadastro:
        cad = ler_tabela(arq)
        col_id = next((c for c in candidatas if c in cad.columns), None)
        col_doc = next((c for c in cad.columns if any(p in c for p in PALAVRAS_DOC)), None)
        if col_id and col_doc:
            partes.append(cad[[col_id, col_doc]].rename(columns={col_id: chave, col_doc: "doc_cadastro"}))
            log(f"  Cadastro {Path(arq).name}: {len(cad)} clientes (cruzando por '{chave}')")
        else:
            log(f"  AVISO: {Path(arq).name} ignorado (precisa de coluna de cliente e de CPF/CNPJ).")
    if not partes:
        return df

    cadastro = pd.concat(partes).dropna().drop_duplicates(subset=[chave])
    df = df.merge(cadastro, on=chave, how="left")
    df["cpf_cnpj"] = df["cpf_cnpj"].fillna(df["doc_cadastro"])
    return df.drop(columns=["doc_cadastro"])


def formatar_documento(valor, mascarar_cpf=False):
    if pd.isna(valor) or not str(valor).strip():
        return None
    d = re.sub(r"\D", "", str(valor))
    if len(d) == 11:
        if mascarar_cpf:
            return f"***.{d[3:6]}.{d[6:9]}-**"
        return f"{d[:3]}.{d[3:6]}.{d[6:9]}-{d[9:]}"
    if len(d) == 14:
        return f"{d[:2]}.{d[2:5]}.{d[5:8]}/{d[8:12]}-{d[12:]}"
    return str(valor).strip()  # tamanho fora do padrão: mantém para conferência


# ------------------------------------------------------------------
# Padronização
# ------------------------------------------------------------------
def _num(x):
    if x is None or (isinstance(x, float) and np.isnan(x)):
        return np.nan
    s = str(x).strip().replace("R$", "").replace(" ", "")
    if not s:
        return np.nan
    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return np.nan


def padronizar(df, arquivos_cadastro=(), mascarar_cpf=False, log=print) -> pd.DataFrame:
    df = extrair_documento(df)
    df = df.rename(columns=RENOMEAR)

    faltando = [c for c in OBRIGATORIAS if c not in df.columns]
    if faltando:
        raise ValueError(
            f"Colunas obrigatórias não encontradas: {faltando}\n"
            f"Colunas que vieram no arquivo: {list(df.columns)}\n"
            "Ajuste o dicionário RENOMEAR no topo do arquivo motor.py."
        )

    ausentes = [c for c in COLUNAS_VALOR if c not in df.columns]
    if ausentes:
        log(f"  AVISO: colunas de valor não encontradas {ausentes}. "
            f"Colunas do arquivo: {list(df.columns)}")
    for col in COLUNAS_VALOR:
        df[col] = df[col].map(_num).round(2) if col in df.columns else np.nan

    # descarta linhas sem ID de transação (totais, rodapés, linhas em branco)
    sem_id = df["id_transacao"].isna() | (df["id_transacao"].astype(str).str.strip() == "")
    if sem_id.any():
        log(f"  {int(sem_id.sum())} linha(s) sem ID de transação descartadas (totais/rodapé).")
        df = df[~sem_id].copy()
    for col in COLUNAS_DATA:
        if col in df.columns:
            df[col] = pd.to_datetime(df[col], errors="coerce", utc=True, dayfirst=True).dt.tz_localize(None)

    df["moeda"] = df["moeda"].str.upper() if "moeda" in df.columns else "N/D"

    df = completar_com_cadastro(df, arquivos_cadastro, log)
    df["cpf_cnpj"] = df["cpf_cnpj"].map(lambda v: formatar_documento(v, mascarar_cpf))

    return df[[c for c in ORDEM if c in df.columns]].copy()


def remover_duplicatas(df):
    antes = len(df)
    repetida = df["id_transacao"].notna() & df.duplicated(subset=["id_transacao"], keep="first")
    df = df[~repetida].copy()
    return df, antes - len(df)


# ------------------------------------------------------------------
# Conferência
# ------------------------------------------------------------------
def resumo_por(df, chaves):
    return (
        df.groupby(chaves, dropna=False)
        .agg(qtd_transacoes=("id_transacao", "count"),
             bruto=("bruto", lambda x: x.sum(min_count=1)),
             tarifas=("tarifa", lambda x: x.sum(min_count=1)),
             liquido=("liquido", lambda x: x.sum(min_count=1)))
        .round(2).reset_index()
    )


def divergencias(df):
    if df[COLUNAS_VALOR].isna().all().all():
        return df.iloc[0:0]
    dif = (df["bruto"] - df["tarifa"] - df["liquido"]).abs()
    return df[(dif > 0.01) | dif.isna()]


# ------------------------------------------------------------------
# Excel
# ------------------------------------------------------------------
def _cabecalho(ws, linha, n):
    for col in range(1, n + 1):
        c = ws.cell(row=linha, column=col)
        c.fill, c.font = COR_CAB, FONTE_CAB
        c.alignment = Alignment(horizontal="center")


def _aba_dados(wb, nome, df):
    ws = wb.create_sheet(nome)
    ws.append(list(df.columns))
    for linha in df.itertuples(index=False):
        ws.append([None if pd.isna(v) else v for v in linha])
    _cabecalho(ws, 1, len(df.columns))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for i, nome_col in enumerate(df.columns, start=1):
        letra = get_column_letter(i)
        amostra = [len(str(c.value)) for c in ws[letra][1:200] if c.value is not None]
        ws.column_dimensions[letra].width = min(max([len(nome_col)] + amostra) + 3, 45)
        fmt = "#,##0.00" if nome_col in COLUNAS_VALOR else (
            "dd/mm/yyyy hh:mm" if nome_col in COLUNAS_DATA else None)
        for c in ws[letra]:
            if c.row > 1:
                c.font = Font(name=FONTE)
                if fmt:
                    c.number_format = fmt


def _bloco(ws, linha, titulo, tabela):
    ws.cell(row=linha, column=1, value=titulo).font = Font(name=FONTE, bold=True, size=12)
    linha += 1
    for j, nome in enumerate(tabela.columns, start=1):
        ws.cell(row=linha, column=j, value=nome)
    _cabecalho(ws, linha, len(tabela.columns))
    for dados in tabela.itertuples(index=False):
        linha += 1
        for j, v in enumerate(dados, start=1):
            c = ws.cell(row=linha, column=j, value=None if pd.isna(v) else v)
            c.font = Font(name=FONTE)
            nome = tabela.columns[j - 1]
            if nome == "qtd_transacoes":
                c.number_format = "#,##0"
            elif nome in ("bruto", "tarifas", "liquido"):
                c.number_format = "#,##0.00"
    return linha + 3


def gerar_excel(df, qtd_lidas, qtd_dup, n_arquivos, pasta_saida: Path) -> Path:
    pasta_saida.mkdir(parents=True, exist_ok=True)
    caminho = pasta_saida / f"conferencia_repasses_{datetime.now():%Y%m%d_%H%M%S}.xlsx"

    df = df.copy()
    tem_repasse = "data_repasse" in df.columns and df["data_repasse"].notna().any()
    col_data = "data_repasse" if tem_repasse else "data_transacao"
    rotulo_mes = "repasse" if tem_repasse else "transação"
    df["mes_repasse"] = (
        df[col_data].dt.strftime("%Y-%m").fillna("N/D") if col_data in df.columns else "N/D"
    )
    div = divergencias(df)

    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Resumo Executivo")
    ws["A1"] = "Resumo Executivo - Conferência de Repasses Stripe"
    ws["A1"].font = Font(name=FONTE, bold=True, size=14)
    ws["A2"] = f"Gerado em {datetime.now():%d/%m/%Y %H:%M}"
    ws["A2"].font = Font(name=FONTE)

    info = [
        ("Arquivos processados", n_arquivos),
        ("Linhas lidas (total dos arquivos)", qtd_lidas),
        ("Duplicatas removidas", qtd_dup),
        ("Transações consolidadas", len(df)),
        ("Transações sem CPF/CNPJ", int(df["cpf_cnpj"].isna().sum())),
        ("Linhas com divergência (bruto - tarifa ≠ líquido)", len(div)),
    ]
    if col_data in df.columns and df[col_data].notna().any():
        info.insert(0, (f"Período ({rotulo_mes})",
                        f"{df[col_data].min():%d/%m/%Y} a {df[col_data].max():%d/%m/%Y}"))
    linha = 4
    for rotulo, valor in info:
        ws.cell(row=linha, column=1, value=rotulo).font = Font(name=FONTE, bold=True)
        ws.cell(row=linha, column=2, value=valor).font = Font(name=FONTE)
        linha += 1

    linha += 1
    linha = _bloco(ws, linha, "Totais por moeda (todas as linhas)", resumo_por(df, ["moeda"]))
    if "categoria" in df.columns:
        eh_payout = df["categoria"].astype(str).str.lower().str.startswith("payout")
        if eh_payout.any():
            linha = _bloco(ws, linha, "Totais por moeda (sem linhas de payout/repasse)",
                           resumo_por(df[~eh_payout], ["moeda"]))
    linha = _bloco(ws, linha, "Totais por arquivo", resumo_por(df, ["arquivo_origem", "moeda"]))
    linha = _bloco(ws, linha, f"Totais por mês de {rotulo_mes}", resumo_por(df, ["mes_repasse", "moeda"]))
    if "categoria" in df.columns:
        linha = _bloco(ws, linha, "Totais por categoria", resumo_por(df, ["categoria", "moeda"]))
    for letra, larg in zip("ABCDEF", [50, 16, 18, 18, 18, 18]):
        ws.column_dimensions[letra].width = larg

    _aba_dados(wb, "Transações", df.drop(columns=["mes_repasse"]))

    wb.save(caminho)
    return caminho


# ------------------------------------------------------------------
# Função principal chamada pela interface
# ------------------------------------------------------------------
def processar(arquivos_repasse, arquivos_cadastro, pasta_saida, mascarar_cpf=False, log=print):
    log("1/4 Lendo arquivos...")
    bruto = ler_repasses(arquivos_repasse, log)
    qtd_lidas, n_arquivos = len(bruto), bruto["arquivo_origem"].nunique()

    log("2/4 Padronizando e buscando CPF/CNPJ...")
    df = padronizar(bruto, arquivos_cadastro, mascarar_cpf, log)
    df, qtd_dup = remover_duplicatas(df)
    log(f"  Duplicatas removidas: {qtd_dup}")

    log("3/4 Conferindo valores...")
    totais = resumo_por(df, ["moeda"])
    n_div = len(divergencias(df))

    log("4/4 Gerando Excel...")
    caminho = gerar_excel(df, qtd_lidas, qtd_dup, n_arquivos, Path(pasta_saida))

    return {
        "caminho": caminho,
        "totais": totais,
        "qtd_lidas": qtd_lidas,
        "qtd_dup": qtd_dup,
        "qtd_final": len(df),
        "sem_documento": int(df["cpf_cnpj"].isna().sum()),
        "divergencias": n_div,
    }